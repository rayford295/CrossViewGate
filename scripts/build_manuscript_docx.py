from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

REPO_ROOT = Path(__file__).resolve().parents[1]

FIGURES = {
    "Figure 1.": [REPO_ROOT / "figures/fig0_pipeline.png"],
    "Figure 2.": [REPO_ROOT / "figures/fig2_oracle_gap.png"],
    "Figure 3.": [REPO_ROOT / "figures/fig3_gate.png"],
    "Figure 4.": [REPO_ROOT / "figures/fig4_fov_intervention.png"],
    "Figure 5.": [REPO_ROOT / "figures/fig5_conflict_density.png"],
    "Figure 6.": [REPO_ROOT / "figures/fig6_qualitative_conflicts.png"],
}


def clean_math(text: str) -> str:
    text = text.replace(
        "$$\\mathrm{closure}(m) = \\frac{\\mathrm{acc}_{\\mathrm{conf}}(m) - \\mathrm{acc}_{\\mathrm{conf}}(\\mathrm{best single view})}{\\mathrm{acc}_{\\mathrm{conf}}(\\mathrm{oracle}) - \\mathrm{acc}_{\\mathrm{conf}}(\\mathrm{best single view})},$$",
        "closure(m) = [acc_conf(m) - acc_conf(best single view)] / [acc_conf(oracle) - acc_conf(best single view)]        (1)",
    )
    text = re.sub(r"\$\\\{x : s\(x\) \\neq r\(x\)\\\}\$", "{x : s(x) != r(x)}", text)
    text = text.replace("$s(x)$", "s(x)").replace("$r(x)$", "r(x)")
    text = text.replace("$m$", "m").replace("$w(x) \\in [0,1]$", "w(x) in [0, 1]")
    text = text.replace("H0: E[delta] = 0", "H0: E[delta] = 0")
    text = re.sub(r"\$([^$]+)\$", r"\1", text)
    return text


def add_runs(paragraph, text: str) -> None:
    text = clean_math(text)
    pos = 0
    for match in re.finditer(r"\*\*(.+?)\*\*|\*(.+?)\*", text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        if match.group(1) is not None:
            run = paragraph.add_run(match.group(1))
            run.bold = True
        else:
            run = paragraph.add_run(match.group(2))
            run.italic = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_table(document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row in enumerate(rows):
        for col_index, cell_text in enumerate(row):
            if col_index >= len(rows[0]):
                continue
            cell = table.cell(row_index, col_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            add_runs(paragraph, cell_text.strip())
            for run in paragraph.runs:
                run.font.size = Pt(9)
                if row_index == 0:
                    run.bold = True
    document.add_paragraph()


def add_figure(document, image_paths: list[Path]) -> None:
    existing = [p for p in image_paths if p.exists()]
    if not existing:
        placeholder = document.add_paragraph()
        run = placeholder.add_run("[Figure placeholder — image file not found]")
        run.italic = True
        placeholder.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    width = Inches(6.0 / len(existing) if len(existing) > 1 else 5.5)
    for path in existing:
        paragraph.add_run().add_picture(str(path), width=width)
        paragraph.add_run("  ")


def main() -> None:
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else REPO_ROOT / "paper/isprs_manuscript.md"
    output = Path(sys.argv[2]) if len(sys.argv) > 2 else REPO_ROOT / "outputs/isprs_manuscript.docx"

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|"):
            table_rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [c.strip() for c in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    table_rows.append(cells)
                index += 1
            if table_rows:
                width = max(len(r) for r in table_rows)
                table_rows = [r + [""] * (width - len(r)) for r in table_rows]
                add_table(document, table_rows)
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped.lstrip("#").strip()
            if level == 1:
                heading = document.add_paragraph()
                run = heading.add_run(clean_math(re.sub(r"\*\*?", "", text)))
                run.bold = True
                run.font.size = Pt(16)
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                document.add_heading(re.sub(r"\*\*?", "", text), level=min(level - 1, 3))
            index += 1
            continue
        for prefix, images in FIGURES.items():
            if stripped.startswith(f"**{prefix}**"):
                add_figure(document, images)
                break
        if stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_runs(paragraph, stripped[2:])
        else:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(paragraph, stripped)
        index += 1

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    print(f"wrote {output}")


if __name__ == "__main__":
    main()
